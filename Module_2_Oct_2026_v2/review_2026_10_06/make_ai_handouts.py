"""Builds the Day 1 AI handouts (Sessions 1 and 3) that go with the printed worksheets.
Texts are read from the R helper files so the handouts and the slides cannot drift apart.
Run from the Module_2_Oct_2026_v2 folder:  python review_2026_10_06/make_ai_handouts.py
"""
import re, sys, shutil
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
LIT = re.compile(r"'((?:[^'\\]|\\.)*)'|(\bhy_table_text\b)")

def r_text(file, var, env=None):
    t = (ROOT / file).read_text(encoding='utf-8')
    i = t.index(var + ' <-')
    j = t.index('\n', t.index('\n', i) + 1) if False else None
    # take the statement: up to the closing of paste0( ... ) at depth 0
    k = t.index('(', i); depth = 0; q = False; e = k
    while e < len(t):
        c = t[e]
        if c == "'" and t[e-1] != '\\': q = not q
        elif not q and c == '(': depth += 1
        elif not q and c == ')':
            depth -= 1
            if depth == 0: break
        e += 1
    stmt = t[k:e]
    out = ''
    for m in LIT.finditer(stmt):
        if m.group(2): out += env[m.group(2)]
        else: out += m.group(1).replace("\\'", "'").replace('\\n', '\n')
    return out

def ascii_minus(s): return s
NAVY = RGBColor(0x06, 0x33, 0x60); GREY = RGBColor(0x54, 0x58, 0x60); LINE = RGBColor(0xB5, 0xBE, 0xC9)

def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.top_margin = s.bottom_margin = Cm(1.1); s.left_margin = s.right_margin = Cm(1.9)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.space_before = Pt(0)
    return d

def para(d, text='', size=11, bold=False, colour=None, after=6, before=0, italic=False, keep=False):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if keep: p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if colour is not None: r.font.color.rgb = colour
    return p

def boxed(d, paras, size=10.5, fill='F3F7FB'):
    t = d.add_table(rows=1, cols=1); t.autofit = True
    c = t.rows[0].cells[0]
    tcPr = c._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}'); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6'); b.set(qn('w:color'), 'A3B7CE'); borders.append(b)
    tcPr.append(borders)
    first = True
    for tx in paras:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(tx); r.font.size = Pt(size)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

def lines(d, n=2):
    for _ in range(n):
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(7)
        r = p.add_run('_' * 78); r.font.size = Pt(11); r.font.color.rgb = LINE

def header(d, sess, title):
    para(d, f'GreenWaste reading workshop | Day 1 | Session {sess} | AI handout', 10, colour=GREY, after=3)
    para(d, title, 16, bold=True, colour=NAVY, after=6)

def datatable(d, rows, widths=(6.5, 3.4, 4.2, 2.4)):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.width = Cm(widths[j])
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
            r = p.add_run(v); r.font.size = Pt(10.5); r.bold = (i == 0)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def split(x): return [p for p in x.split('\n\n') if p.strip()]
NOTICE = 'GreenWaste is a fictional training case.'
WRITTEN = 'This response was written for the exercise. It is not a real AI answer.'

def session1():
    env = {'hy_table_text': r_text('session1_hybrid.R', 'hy_table_text')}
    ai = r_text('session1_review.R', 'review_ai'); fb = r_text('session1_review.R', 'review_fallback')
    pr = r_text('session1_hybrid.R', 'hy_prompt', env)
    d = new_doc()
    header(d, 1, 'AI exercise: read the regression table with an AI')
    para(d, 'You will mark a written AI response, then run an improved prompt in Microsoft Copilot and compare the two. Keep page 1 of your worksheet (the regression table) and your AI good-practice card beside you. The whole exercise takes about 15 minutes.', 11)
    para(d, NOTICE, 9, colour=GREY, after=8)
    para(d, 'The table you are reading (a copy of the table on your worksheet)', 11, bold=True, after=3, keep=True)
    datatable(d, [['Result row', 'Coefficient AED', '95% interval AED', 'p-value'],
                  ['Before GreenWaste', '1,432', '1,418 to 1,446', '< 0.001'],
                  ['After minus before', '-669', '-684 to -654', '< 0.001']])
    para(d, 'Part A: Mark the response (about 8 minutes)', 13, bold=True, colour=NAVY, after=4, before=4, keep=True)
    para(d, 'Naive prompt (a prompt with no context, as a busy person might type it): “Explain this table and advise whether to expand GreenWaste.” ' + WRITTEN, 10, colour=GREY, after=6)
    boxed(d, split(ai))
    para(d, 'Underline the claims that the regression table supports. Cross out the claims that it does not support. Use the table, not how convincing the paragraph sounds.', 11, bold=True, after=6)
    para(d, 'Which claim would you correct first, and what would you write instead?', 11, after=4)
    lines(d, 2)
    page_break(d)
    header(d, 1, 'Part B: Run the improved prompt, then compare')
    para(d, 'Copy the prompt below into Microsoft Copilot. The table is included in the prompt. If you are not using Copilot, read the prepared response below after you finish Part A.', 11)
    para(d, 'Improved prompt', 11, bold=True, after=3, keep=True)
    boxed(d, split(pr), size=10, fill='EEF4FA')
    para(d, 'After the AI has answered, write your comparison.', 11, bold=True, after=4)
    for q in ['(a) Which question that the AI asked helped you most?',
              '(b) What improved compared with the response in Part A?',
              '(c) Which claim still needs checking against the table?']:
        para(d, q, 11, after=2); lines(d, 1)
    para(d, 'A better prompt reduces mistakes. It does not remove the need to check every number and claim against the table.', 10, colour=GREY, italic=True, after=8)
    para(d, 'Prepared response for participants not using Copilot', 13, bold=True, colour=NAVY, after=4, before=6, keep=True)
    para(d, 'Read this after you have marked the response in Part A. It is an example of a better answer. An AI tool may answer differently. ' + WRITTEN, 10.5, colour=GREY)
    boxed(d, split(fb), size=10)
    para(d, '(a) Which sentences does the regression table support?', 10.5, after=1); lines(d, 1)
    para(d, '(b) What additional evidence would answer the clarifying question at the start of the response?', 10.5, after=2); lines(d, 1)
    out = ROOT / 'Oct12_session1_ai_handout.docx'; d.save(out); return out

def session3():
    ai = r_text('session3_hybrid.R', 'hy3_ai') if False else None
    pr = r_text('session3_review.R', 's3_prompt').replace('Cite H3-A.', 'Cite the pilot trial information.').replace('Use only H3-A:', 'Use only this pilot trial information:')
    fb = r_text('session3_review.R', 's3_fallback').replace('H3-A reports', 'The pilot trial reports')
    assert 'H3-A' not in pr + fb
    d = new_doc()
    header(d, 3, 'AI exercise: advice for the minister')
    para(d, 'On page 2 of your worksheet you checked an illustrative AI draft. Now run an improved prompt in Microsoft Copilot and compare the advice with the draft you marked and with your own decision note. Keep page 1 of your worksheet (the pilot trial) and your AI good-practice card beside you. This exercise takes about 7 minutes.', 11)
    para(d, NOTICE, 9, colour=GREY, after=8)
    para(d, 'The pilot trial result (a copy from your worksheet)', 11, bold=True, after=3, keep=True)
    para(d, 'In one district, a lottery assigned 400 businesses to join or wait. The trial compared their annual waste costs 12 months later.', 10.5, after=4)
    datatable(d, [['Result row', 'Coefficient AED', '95% interval AED', 'p-value'],
                  ['Took part (lottery)', '-1,014', '-1,167 to -862', '< 0.001']])
    para(d, 'Decision rule: annual saving of at least 1,000 AED per business. Negative means lower costs.', 10.5, after=8)
    para(d, 'Run the improved prompt, then compare', 13, bold=True, colour=NAVY, after=4, before=2, keep=True)
    para(d, 'Copy the prompt below into Microsoft Copilot. It contains the pilot facts, so you can type or paste all of it. If you are not using Copilot, read the prepared response on the next page after you finish.', 11)
    para(d, 'Improved prompt', 11, bold=True, after=3, keep=True)
    boxed(d, split(pr), fill='EEF4FA')
    for q in ['(a) Which question that the AI asked helped you most?',
              '(b) What improved compared with the draft you marked and with your own decision note?',
              '(c) Which claim still needs checking against the pilot trial result on page 1 of your worksheet?']:
        para(d, q, 11, after=2); lines(d, 1)
    para(d, 'A better prompt reduces mistakes. It does not remove the need to check every number and claim against the pilot trial result.', 10, colour=GREY, italic=True)
    page_break(d)
    header(d, 3, 'Prepared response for participants not using Copilot')
    para(d, 'Read this after you have marked the draft on page 2 of your worksheet. It is an example of a better answer. An AI tool may answer differently. ' + WRITTEN, 10.5, colour=GREY)
    boxed(d, split(fb))
    for q in ['(a) Which clarifying question in the response helps most?',
              '(b) Does the advice separate some reduction in cost, our rule of at least a 1,000 AED reduction, wider reach and full programme costs?',
              '(c) Which claim would you check or strengthen before sharing this advice?']:
        para(d, q, 11, after=2); lines(d, 2)
    out = ROOT / 'Oct12_session3_ai_handout.docx'; d.save(out); return out

if __name__ == '__main__':
    for f in (session1(), session3()):
        print('wrote', f)
