# Build a PDF and a PowerPoint per deck from the captured slide images.
# The PowerPoint has one full-slide picture per slide and the speaker notes in
# the notes pane; it is for sharing, not for editing.
import json, os, sys
import fitz
from pptx import Presentation
from pptx.util import Emu

CAP = 'cap'
OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)
W_EMU, H_EMU = 12192000, 6858000          # 13.333 x 7.5 in, 16:9

for deck in sorted(os.listdir(CAP)):
    folder = os.path.join(CAP, deck)
    meta_file = os.path.join(folder, 'slides.json')
    if not os.path.exists(meta_file):
        print('SKIP (no capture)', deck); continue
    meta = json.load(open(meta_file, encoding='utf8'))
    name = deck.replace('_live', '')
    title = meta[0]['title'] if meta else name

    pdf = fitz.open()
    for m in meta:
        img = os.path.join(folder, m['file'])
        page = pdf.new_page(width=960, height=540)
        page.insert_image(page.rect, filename=img)
    pdf.set_metadata({'title': title, 'author': '3ie', 'subject': 'Module 2, Abu Dhabi, October 2026'})
    pdf.save(os.path.join(OUT, name + '.pdf'), deflate=True)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W_EMU), Emu(H_EMU)
    blank = prs.slide_layouts[6]
    for m in meta:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(os.path.join(folder, m['file']), 0, 0, width=Emu(W_EMU), height=Emu(H_EMU))
        if m.get('notes'):
            s.notes_slide.notes_text_frame.text = m['notes']
    prs.core_properties.title = title
    prs.core_properties.author = '3ie'
    prs.save(os.path.join(OUT, name + '.pptx'))
    print(f'{name}: {len(meta)} slides, pdf {os.path.getsize(os.path.join(OUT, name + ".pdf"))//1024} KB, '
          f'pptx {os.path.getsize(os.path.join(OUT, name + ".pptx"))//1024} KB')
